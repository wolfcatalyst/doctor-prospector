"""Excel export with bidirectional hyperlinks between sheets.

Generates a multi-sheet Excel workbook where:
- Physicians sheet links forward to Trials, Publications, Payments, Contact sheets
- Detail sheets link back to the physician's row on the Physicians sheet
- External links go to PubMed, NPI Registry, Doximity, Google
"""

import logging
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .models import ExportData, ExportOptions, HospitalExportRecord, PhysicianExportRecord
from .styles import (
    HEADER_ALIGNMENT,
    HEADER_FILL,
    HEADER_FONT,
    LINK_FONT,
    MISSING_EXCEL,
    MISSING_FONT,
    TAB_COLORS,
)

logger = logging.getLogger(__name__)


def export_excel(data: ExportData) -> bytes:
    """Generate Excel workbook with bidirectional hyperlinks.

    Returns:
        Bytes of the .xlsx file.
    """
    wb = Workbook()
    options = data.options
    physicians = data.all_physicians

    # Pre-compute row mappings for bidirectional linking
    row_map = _build_row_mappings(physicians, data.hospitals, options)

    # Physicians sheet is always created (it's the default active sheet)
    ws_physicians = wb.active
    ws_physicians.title = "Physicians"
    _write_physicians_sheet(ws_physicians, physicians, row_map, options)

    if options.include_hospitals:
        ws_hospitals = wb.create_sheet("Hospitals")
        _write_hospitals_sheet(ws_hospitals, data.hospitals, row_map)

    if options.include_trials:
        ws_trials = wb.create_sheet("Clinical Trials")
        _write_trials_sheet(ws_trials, physicians, row_map)

    if options.include_publications:
        ws_pubs = wb.create_sheet("Publications")
        _write_publications_sheet(ws_pubs, physicians, row_map)

    if options.include_payments:
        ws_payments = wb.create_sheet("Device Payments")
        _write_payments_sheet(ws_payments, physicians, row_map)

    if options.include_contact:
        ws_contact = wb.create_sheet("Contact Info")
        _write_contact_sheet(ws_contact, physicians, row_map)

    # Apply visual polish to all sheets
    for ws in wb.worksheets:
        _apply_sheet_styling(ws)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _build_row_mappings(
    physicians: list[PhysicianExportRecord],
    hospitals: list[HospitalExportRecord],
    options: ExportOptions,
) -> dict[str, Any]:
    """Pre-compute row positions for every physician across all sheets.

    Must be computed before any sheet is written, because links are bidirectional.
    """
    mapping: dict[str, Any] = {
        "physicians": {},
        "trials": {},
        "publications": {},
        "payments": {},
        "contact": {},
        "hospitals": {},
        "hospital_first_physician": {},
    }

    # Physicians sheet: row 1 = header, row 2+ = data (1-indexed)
    for idx in range(len(physicians)):
        mapping["physicians"][idx] = idx + 2

    # Trials sheet: each physician can have multiple trials
    trial_row = 2
    for idx, p in enumerate(physicians):
        rows = []
        if options.include_trials:
            for _ in p.trials:
                rows.append(trial_row)
                trial_row += 1
        mapping["trials"][idx] = rows

    # Publications sheet: each physician can have multiple publications
    pub_row = 2
    for idx, p in enumerate(physicians):
        rows = []
        if options.include_publications:
            for _ in p.publications:
                rows.append(pub_row)
                pub_row += 1
        mapping["publications"][idx] = rows

    # Payments sheet: one row per company per physician
    pay_row = 2
    for idx, p in enumerate(physicians):
        rows = []
        if options.include_payments and p.payments and p.payments.get("companies"):
            for _ in p.payments["companies"]:
                rows.append(pay_row)
                pay_row += 1
        mapping["payments"][idx] = rows

    # Contact sheet: 1:1 with physicians
    for idx in range(len(physicians)):
        mapping["contact"][idx] = idx + 2

    # Hospitals sheet: one row per unique hospital
    seen_hospitals: dict[str, int] = {}
    hospital_row = 2
    for idx, p in enumerate(physicians):
        if p.hospital_name and p.hospital_name not in seen_hospitals:
            seen_hospitals[p.hospital_name] = hospital_row
            mapping["hospital_first_physician"][p.hospital_name] = idx
            hospital_row += 1
    mapping["hospitals"] = seen_hospitals

    # Also add hospitals that have no physicians (from the hospitals list)
    for h in hospitals:
        if h.name and h.name not in seen_hospitals:
            seen_hospitals[h.name] = hospital_row
            hospital_row += 1

    return mapping


def _write_header(ws, columns: list[str]) -> None:
    """Write styled header row."""
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT


def _write_cell(ws, row: int, col: int, value: Any, missing: bool = False):
    """Write a cell, applying missing-data styling if needed."""
    display = MISSING_EXCEL if missing or value is None or value == "" else value
    cell = ws.cell(row=row, column=col, value=display)
    if missing or value is None or value == "":
        cell.font = MISSING_FONT
    return cell


def _write_link_cell(ws, row: int, col: int, value: str, hyperlink: str):
    """Write a cell with a hyperlink."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.hyperlink = hyperlink
    cell.font = LINK_FONT
    return cell


def _fmt_missing(value: Any) -> bool:
    """Check if a value should be treated as missing."""
    return value is None or value == "" or value == []


# ---------------------------------------------------------------------------
# Sheet Writers
# ---------------------------------------------------------------------------


def _write_physicians_sheet(
    ws, physicians: list[PhysicianExportRecord], row_map: dict, options: ExportOptions
) -> None:
    """Write the main Physicians index sheet with forward links."""
    # Build column list dynamically
    columns = ["Name", "Hospital", "Specialty", "Volume"]
    if options.include_historical:
        columns.append("Trend")
    if options.include_trials:
        columns.append("Trials")
    if options.include_publications:
        columns.append("Pubs")
    if options.include_payments:
        columns.append("Payments")
    if options.include_contact:
        columns.append("Contact")

    _write_header(ws, columns)

    for p_idx, physician in enumerate(physicians):
        row = row_map["physicians"][p_idx]
        col = 1

        # Name
        ws.cell(row=row, column=col, value=physician.full_name)
        col += 1

        # Hospital (link to Hospitals sheet if available)
        if options.include_hospitals and physician.hospital_name in row_map["hospitals"]:
            h_row = row_map["hospitals"][physician.hospital_name]
            _write_link_cell(ws, row, col, physician.hospital_name, f"#'Hospitals'!A{h_row}")
        else:
            _write_cell(ws, row, col, physician.hospital_name)
        col += 1

        # Specialty
        _write_cell(ws, row, col, physician.primary_specialty, missing=_fmt_missing(physician.primary_specialty))
        col += 1

        # Volume
        _write_cell(ws, row, col, physician.tavr_volume if physician.tavr_volume else None, missing=not physician.tavr_volume)
        col += 1

        # Trend (if included)
        if options.include_historical:
            if physician.historical:
                growth = physician.historical.get("avg_annual_growth_pct", 0)
                indicator = physician.historical.get("trend_indicator", "")
                sign = "+" if growth > 0 else ""
                ws.cell(row=row, column=col, value=f"{sign}{growth:.1f}% {indicator}")
            else:
                _write_cell(ws, row, col, None, missing=True)
            col += 1

        # Trials count with forward link
        if options.include_trials:
            trial_rows = row_map["trials"].get(p_idx, [])
            if trial_rows:
                _write_link_cell(
                    ws, row, col,
                    f"{len(trial_rows)}",
                    f"#'Clinical Trials'!A{trial_rows[0]}",
                )
            else:
                _write_cell(ws, row, col, MISSING_EXCEL, missing=True)
            col += 1

        # Publications count with forward link
        if options.include_publications:
            pub_rows = row_map["publications"].get(p_idx, [])
            if pub_rows:
                _write_link_cell(
                    ws, row, col,
                    f"{len(pub_rows)}",
                    f"#'Publications'!A{pub_rows[0]}",
                )
            else:
                _write_cell(ws, row, col, MISSING_EXCEL, missing=True)
            col += 1

        # Payments total with forward link
        if options.include_payments:
            pay_rows = row_map["payments"].get(p_idx, [])
            if pay_rows and physician.payments:
                total = physician.payments.get("total_payments", 0)
                _write_link_cell(
                    ws, row, col,
                    f"${total:,.0f}",
                    f"#'Device Payments'!A{pay_rows[0]}",
                )
            else:
                _write_cell(ws, row, col, MISSING_EXCEL, missing=True)
            col += 1

        # Contact link
        if options.include_contact:
            contact_row = row_map["contact"][p_idx]
            _write_link_cell(ws, row, col, "View", f"#'Contact Info'!A{contact_row}")
            col += 1


def _write_hospitals_sheet(
    ws, hospitals: list[HospitalExportRecord], row_map: dict
) -> None:
    """Write Hospitals sheet with link to first physician."""
    columns = [
        "Hospital", "Location", "Annual Volume", "Star Rating", "Program Since",
        "Beds", "Cath Lab", "Open Heart", "Teaching", "Ownership", "Physicians",
    ]
    _write_header(ws, columns)

    for hospital in hospitals:
        h_row = row_map["hospitals"].get(hospital.name)
        if h_row is None:
            continue

        ws.cell(row=h_row, column=1, value=hospital.name)
        ws.cell(row=h_row, column=2, value=f"{hospital.city}, {hospital.state}")
        ws.cell(row=h_row, column=3, value=hospital.annual_tavr_volume)
        _write_cell(ws, h_row, 4, hospital.star_rating, missing=_fmt_missing(hospital.star_rating))
        _write_cell(ws, h_row, 5, hospital.first_tavr_date, missing=_fmt_missing(hospital.first_tavr_date))
        # POS facility details
        _write_cell(ws, h_row, 6, hospital.bed_count, missing=_fmt_missing(hospital.bed_count))
        _write_cell(ws, h_row, 7,
                     "Yes" if hospital.has_cath_lab else ("No" if hospital.has_cath_lab is not None else None),
                     missing=_fmt_missing(hospital.has_cath_lab))
        _write_cell(ws, h_row, 8,
                     "Yes" if hospital.has_open_heart_surgery else ("No" if hospital.has_open_heart_surgery is not None else None),
                     missing=_fmt_missing(hospital.has_open_heart_surgery))
        _write_cell(ws, h_row, 9, hospital.teaching_status, missing=_fmt_missing(hospital.teaching_status))
        _write_cell(ws, h_row, 10, hospital.ownership_type, missing=_fmt_missing(hospital.ownership_type))

        # Link to first physician for this hospital
        first_p_idx = row_map["hospital_first_physician"].get(hospital.name)
        if first_p_idx is not None:
            phys_row = row_map["physicians"][first_p_idx]
            _write_link_cell(
                ws, h_row, 11,
                f"{len(hospital.physicians)}",
                f"#'Physicians'!A{phys_row}",
            )
        else:
            ws.cell(row=h_row, column=11, value="0")


def _write_trials_sheet(
    ws, physicians: list[PhysicianExportRecord], row_map: dict
) -> None:
    """Write Clinical Trials sheet with back-links to Physicians."""
    columns = ["Physician", "Hospital", "Trial Name", "NCT ID", "Sponsor", "Role", "Status"]
    _write_header(ws, columns)

    for p_idx, physician in enumerate(physicians):
        trial_rows = row_map["trials"].get(p_idx, [])
        phys_row = row_map["physicians"][p_idx]

        for t_idx, trial in enumerate(physician.trials):
            if t_idx >= len(trial_rows):
                break
            row = trial_rows[t_idx]

            # Back-link to physician
            _write_link_cell(ws, row, 1, physician.full_name, f"#'Physicians'!A{phys_row}")

            ws.cell(row=row, column=2, value=physician.hospital_name)
            _write_cell(ws, row, 3, trial.get("title"), missing=_fmt_missing(trial.get("title")))
            _write_cell(ws, row, 4, trial.get("nct_id"), missing=_fmt_missing(trial.get("nct_id")))
            _write_cell(ws, row, 5, trial.get("lead_sponsor"), missing=_fmt_missing(trial.get("lead_sponsor")))

            role = _find_physician_role(physician, trial)
            _write_cell(ws, row, 6, role, missing=_fmt_missing(role))

            _write_cell(ws, row, 7, trial.get("overall_status"), missing=_fmt_missing(trial.get("overall_status")))


def _write_publications_sheet(
    ws, physicians: list[PhysicianExportRecord], row_map: dict
) -> None:
    """Write Publications sheet with back-links and external PubMed links."""
    columns = ["Physician", "Hospital", "Title", "Journal", "Year", "ORCID Verified", "PubMed"]
    _write_header(ws, columns)

    for p_idx, physician in enumerate(physicians):
        pub_rows = row_map["publications"].get(p_idx, [])
        phys_row = row_map["physicians"][p_idx]

        for pub_idx, pub in enumerate(physician.publications):
            if pub_idx >= len(pub_rows):
                break
            row = pub_rows[pub_idx]

            # Back-link to physician
            _write_link_cell(ws, row, 1, physician.full_name, f"#'Physicians'!A{phys_row}")

            ws.cell(row=row, column=2, value=physician.hospital_name)
            _write_cell(ws, row, 3, pub.get("title"), missing=_fmt_missing(pub.get("title")))
            _write_cell(ws, row, 4, pub.get("journal"), missing=_fmt_missing(pub.get("journal")))

            # Extract year from pub_date
            pub_date = pub.get("pub_date", "")
            year = pub_date[:4] if pub_date and len(str(pub_date)) >= 4 else None
            _write_cell(ws, row, 5, year, missing=_fmt_missing(year))

            # ORCID verification badge
            badge = pub.get("confidence_badge", "")
            _write_cell(ws, row, 6, badge, missing=_fmt_missing(badge))

            # External PubMed link
            pubmed_url = pub.get("pubmed_url", "")
            if pubmed_url:
                _write_link_cell(ws, row, 7, "PubMed", pubmed_url)
            else:
                _write_cell(ws, row, 7, None, missing=True)


def _write_payments_sheet(
    ws, physicians: list[PhysicianExportRecord], row_map: dict
) -> None:
    """Write Device Payments sheet - one row per company per physician."""
    columns = ["Physician", "Hospital", "Company", "Amount", "Year Range", "Payment Types"]
    _write_header(ws, columns)

    for p_idx, physician in enumerate(physicians):
        pay_rows = row_map["payments"].get(p_idx, [])
        phys_row = row_map["physicians"][p_idx]

        if not physician.payments or not physician.payments.get("companies"):
            continue

        companies = physician.payments.get("companies", {})
        companies_sorted = sorted(companies.items(), key=lambda x: x[1], reverse=True)

        for c_idx, (company, amount) in enumerate(companies_sorted):
            if c_idx >= len(pay_rows):
                break
            row = pay_rows[c_idx]

            # Back-link to physician
            _write_link_cell(ws, row, 1, physician.full_name, f"#'Physicians'!A{phys_row}")

            ws.cell(row=row, column=2, value=physician.hospital_name)
            ws.cell(row=row, column=3, value=company)
            ws.cell(row=row, column=4, value=f"${amount:,.0f}" if amount else MISSING_EXCEL)

            # Year range
            years = physician.payments.get("years", {})
            if years:
                year_keys = [int(y) for y in years.keys()]
                year_range = f"{min(year_keys)}-{max(year_keys)}" if len(year_keys) > 1 else str(year_keys[0])
                ws.cell(row=row, column=5, value=year_range)
            else:
                _write_cell(ws, row, 5, None, missing=True)

            # Top payment types
            ptypes = physician.payments.get("payment_types", {})
            if ptypes:
                top_types = sorted(ptypes.items(), key=lambda x: x[1], reverse=True)[:3]
                type_str = ", ".join(t[0] for t in top_types)
                ws.cell(row=row, column=6, value=type_str)
            else:
                _write_cell(ws, row, 6, None, missing=True)


def _write_contact_sheet(
    ws, physicians: list[PhysicianExportRecord], row_map: dict
) -> None:
    """Write Contact Info sheet with back-links and external URLs."""
    columns = ["Physician", "Hospital", "Practice Address", "Phone", "NPI Registry", "Doximity", "Google"]
    _write_header(ws, columns)

    for p_idx, physician in enumerate(physicians):
        row = row_map["contact"][p_idx]
        phys_row = row_map["physicians"][p_idx]

        # Back-link to physician
        _write_link_cell(ws, row, 1, physician.full_name, f"#'Physicians'!A{phys_row}")

        ws.cell(row=row, column=2, value=physician.hospital_name)

        # Assemble address
        address = _build_address(physician)
        _write_cell(ws, row, 3, address, missing=_fmt_missing(address))

        # Phone
        phone = _format_phone(physician.practice_phone)
        _write_cell(ws, row, 4, phone, missing=_fmt_missing(phone))

        # External links
        for col, url, label in [
            (5, physician.npi_registry_url, "NPI"),
            (6, physician.doximity_url, "Doximity"),
            (7, physician.google_search_url, "Google"),
        ]:
            if url:
                _write_link_cell(ws, row, col, label, url)
            else:
                _write_cell(ws, row, col, None, missing=True)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _apply_sheet_styling(ws) -> None:
    """Apply visual polish to a worksheet."""
    # Freeze header row
    ws.freeze_panes = "A2"

    # Sheet tab color
    tab_color = TAB_COLORS.get(ws.title)
    if tab_color:
        ws.sheet_properties.tabColor = tab_color

    # Auto-fit column widths (approximate)
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                cell_len = len(str(cell.value))
                max_length = max(max_length, cell_len)
        adjusted_width = min(max(max_length + 2, 10), 50)
        ws.column_dimensions[col_letter].width = adjusted_width


def _find_physician_role(physician: PhysicianExportRecord, trial: dict) -> str | None:
    """Find a physician's role in a specific clinical trial by name matching."""
    if not physician.full_name:
        return None

    # Try matching on last name
    name_parts = physician.full_name.replace(",", "").split()
    last_name = name_parts[-1].lower() if name_parts else ""

    for inv in trial.get("investigators", []):
        inv_name = (inv.get("name", "") or "").lower()
        if last_name and last_name in inv_name:
            return inv.get("role", "Investigator")

    return None


def _build_address(physician: PhysicianExportRecord) -> str | None:
    """Build a formatted address string from physician fields."""
    parts = []
    if physician.practice_address_line1:
        parts.append(physician.practice_address_line1)
    if physician.practice_address_line2:
        parts.append(physician.practice_address_line2)

    city_state_zip = []
    if physician.practice_city:
        city_state_zip.append(physician.practice_city)
    if physician.practice_state:
        city_state_zip.append(physician.practice_state)
    if physician.practice_postal_code:
        city_state_zip.append(physician.practice_postal_code)

    if city_state_zip:
        parts.append(", ".join(city_state_zip[:2]) + (" " + city_state_zip[2] if len(city_state_zip) > 2 else ""))

    if parts:
        return ", ".join(parts)

    # Fallback to practice_address
    return physician.practice_address


def _format_phone(phone: str | None) -> str | None:
    """Format a phone number as (XXX) XXX-XXXX."""
    if not phone:
        return None
    digits = "".join(c for c in phone if c.isdigit())
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    if len(digits) == 11 and digits[0] == "1":
        return f"({digits[1:4]}) {digits[4:7]}-{digits[7:]}"
    return phone
